import json
from pathlib import Path
import sqlite3
import threading

import numpy as np
import pytest
from tokenizers import Tokenizer

from echomind.knowledge import KnowledgeError, KnowledgeService
from echomind.knowledge_import import DocumentPreview, Section, chunk_sections, preview_document
from echomind.knowledge_model import MODEL_ID


@pytest.fixture
def store(tmp_path):
    return KnowledgeService(tmp_path / 'knowledge')


def test_manual_question_edit_delete_and_short_aliases(store):
    identifier = store.save_question('JavaScript 事件循环', '先同步代码，再微任务，然后下一项任务。', 'JS, event loop', '前端')
    assert store.retrieve('JS')[0].document_id == identifier
    assert store.retrieve('请解释一下JavaScript事件循环')[0].source == '手动题目'
    assert not store.retrieve('今晚吃什么')
    store.save_question('Promise 状态', 'pending fulfilled rejected', document_id=identifier)
    assert not store.retrieve('JS')
    assert store.retrieve('Promise 状态')[0].text == 'pending fulfilled rejected'
    store.delete(identifier)
    assert not store.list_documents() and not store.retrieve('Promise')


def test_trigram_short_tags_and_sql_input_are_data(store):
    store.save_question('异步编程基础', '通过事件循环安排异步回调。', tags='JS')
    assert store.retrieve('JS')
    assert not store.retrieve('" OR 1=1; DROP TABLE documents; --')
    assert len(store.list_documents()) == 1
    assert store.retrieve('异步回调')[0].text == '通过事件循环安排异步回调。'


def test_preview_import_deduplicates_and_replaces_atomically(store, tmp_path):
    path = tmp_path / '题库.md'
    path.write_text('# Redis\n\nRedis 基于内存。\n\n# Hooks\n\n函数组件管理状态。', encoding='utf-8')
    preview = preview_document(path)
    assert preview.sections[1].title == 'Redis'
    assert '行' in preview.sections[1].location
    identifier = store.import_document(preview)
    assert store.import_document(preview) == identifier
    assert len(store.list_documents()) == 1
    path.write_text('# Promise\n\nPromise 处理异步结果。', encoding='utf-8')
    assert store.import_document(preview_document(path)) == identifier
    assert not store.retrieve('Redis')
    assert store.retrieve('Promise')
    store.delete(identifier)
    assert path.is_file()


def test_cancel_preserves_existing_version(store, tmp_path):
    path = tmp_path / '题库.txt'
    path.write_text('Redis 在内存读写。', encoding='utf-8')
    identifier = store.import_document(preview_document(path))
    cancel = threading.Event()
    cancel.set()
    path.write_text('完全替换原资料。', encoding='utf-8')
    with pytest.raises(KnowledgeError, match='取消'):
        store.import_document(preview_document(path), cancel=cancel)
    assert store.document(identifier)['sections']
    assert store.retrieve('题库')


def test_busy_database_keeps_old_data(store):
    identifier = store.save_question('Redis', '内存数据库')
    with store.connect() as db:
        db.execute('BEGIN IMMEDIATE')
        with pytest.raises(KnowledgeError, match='稍后'):
            store.save_question('替换', '新数据', document_id=identifier)
    assert store.document(identifier)['question'] == 'Redis'


def test_chunk_token_budget_overlap_and_source():
    title = 'React Hooks' * 10
    source = '这是面试知识，介绍 React useState 和 useEffect。' * 100
    pieces = chunk_sections([Section(title, '第 2 页', source)])
    tokenizer = Tokenizer.from_file(str(Path(__file__).parents[1] / 'echomind/knowledge_tokenizer.json'))
    assert len(pieces) > 1
    for piece in pieces:
        assert piece.location == '第 2 页'
        assert len(tokenizer.encode(piece.title + '\n' + piece.text).ids) <= 384
    assert pieces[0].text[-20:] in pieces[1].text


def test_docx_and_pdf_sources_and_invalid_documents(tmp_path):
    from docx import Document
    from pypdf import PdfWriter
    doc = Document()
    doc.add_heading('Hooks', 1)
    doc.add_paragraph('函数组件可以使用状态。')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'useState'
    table.cell(0, 1).text = '状态'
    path = tmp_path / '题库.docx'
    doc.save(path)
    result = preview_document(path)
    assert result.sections[-1].title == 'Hooks'
    assert 'useState' in result.sections[-1].text
    assert '段落/表格' in result.sections[-1].location
    pdf = tmp_path / '扫描件.pdf'
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(pdf)
    with pytest.raises(KnowledgeError, match='OCR'):
        preview_document(pdf)
    path.write_bytes(b'corrupt')
    with pytest.raises(KnowledgeError, match='损坏'):
        preview_document(path)
    with pytest.raises(KnowledgeError, match='DOC'):
        preview_document(tmp_path / '旧题库.doc')


class FakeEmbedder:
    def encode(self, texts):
        result = np.zeros((len(texts), 512), dtype=np.float32)
        for index, text in enumerate(texts):
            result[index, 0 if '缓存' in text or 'Redis' in text else 1] = 1
        return result


def test_semantic_retrieval_model_versions_and_vector_validation(store):
    identifier = store.save_question('Redis 速度', '内存读写')
    store._embedder = FakeEmbedder()
    store.prepare()
    assert store.retrieve('缓存为什么性能好')[0].document_id == identifier
    assert not store.retrieve('其他问题')
    with store.connect() as db:
        db.execute("UPDATE chunks SET model='wrong-model'")
        db.execute("UPDATE meta SET value=value+1 WHERE key='revision'")
    assert not store.retrieve('缓存性能')
    with pytest.raises(KnowledgeError, match='无效向量'):
        store._vector([float('nan')] * 512)


def test_snapshot_survives_source_deletion_and_indexes_do_not_reuse_ids(store):
    identifier = store.save_question('React Hooks', '原始答案')
    hit = store.retrieve('React Hooks')[0]
    store.delete(identifier)
    second = store.save_question('Redis', '新答案')
    assert hit.text == '原始答案'
    assert store.retrieve('Redis')[0].chunk_id > hit.chunk_id
    assert second != identifier


def test_metadata_only_document_edit_reindexes_and_survives_reopen(store):
    preview = DocumentPreview('旧标题', '题库.md', 'fixture.md', 'digest', (Section('旧标题', '第 1 行', '资料正文'),))
    identifier = store.import_document(preview)
    changed = DocumentPreview('新标题', preview.filename, preview.origin, preview.digest, (Section('新标题', '手动修订正文', '资料正文'),))
    store._save(changed, document_id=identifier, aliases='JS', tags='技术')
    reopened = KnowledgeService(store.directory)
    assert reopened.document(identifier)['title'] == '新标题'
    assert reopened.retrieve('JS')[0].document_id == identifier
    assert not reopened.retrieve('旧标题')


def test_headings_without_blank_lines_and_question_boundaries(tmp_path):
    path = tmp_path / '题库.md'
    path.write_text('# Redis\n内存读写。\n# Hooks\n函数状态。', encoding='utf-8')
    preview = preview_document(path)
    assert [section.title for section in preview.sections] == ['Redis', 'Hooks']
    assert [section.location for section in preview.sections] == ['第 1 行', '第 3 行']
    pieces = chunk_sections([Section('面试题', '第 2 页', '问：Redis 为什么快\n答：内存读写\n问：Promise 有哪些状态\n答：三种状态')])
    assert len(pieces) == 2
    assert 'Promise' not in pieces[0].text
    assert pieces[1].location == '第 2 页'


def test_text_pdf_keeps_actual_page_number(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 10 100 Td (Redis interview question) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    path = tmp_path / 'text.pdf'
    writer.write(path)
    preview = preview_document(path)
    assert preview.sections[0].location == '第 2 页'
    assert 'Redis' in preview.sections[0].text
